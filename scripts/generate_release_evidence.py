import os
import logging
import requests

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
# -----------------------------------------------------------------------------
# Logging
# -----------------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
)

logger = logging.getLogger(__name__)

# -----------------------------------------------------------------------------
# GitHub Configuration
# -----------------------------------------------------------------------------
GITHUB_TOKEN = os.getenv("GITHUB_TOKEN")
GITHUB_API_URL = os.getenv("GITHUB_API_URL", "https://api.github.com")
GITHUB_SERVER_URL = os.getenv("GITHUB_SERVER_URL", "https://github.com")

GITHUB_REPOSITORY = os.getenv("GITHUB_REPOSITORY")
GITHUB_RUN_ID = os.getenv("GITHUB_RUN_ID")
GITHUB_SHA = os.getenv("GITHUB_SHA")
GITHUB_REF_NAME = os.getenv("GITHUB_REF_NAME")
GITHUB_ACTOR = os.getenv("GITHUB_ACTOR")
GITHUB_WORKFLOW = os.getenv("GITHUB_WORKFLOW")
GITHUB_RUN_NUMBER = os.getenv("GITHUB_RUN_NUMBER")

# -----------------------------------------------------------------------------
# GitHub REST API Helper
# -----------------------------------------------------------------------------
def github_get(endpoint: str, params: dict | None = None) -> dict:
    """
    Executes a GET request against the GitHub REST API.
    """

    url = f"{GITHUB_API_URL}{endpoint}"

    headers = {
        "Authorization": f"Bearer {GITHUB_TOKEN}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
        "User-Agent": "ReleaseEvidenceGenerator",
    }

    logger.info("GET %s", url)

    response = requests.get(
        url=url,
        headers=headers,
        params=params,
        timeout=30,
    )

    response.raise_for_status()

    return response.json()
# Create reports directory
os.makedirs("reports", exist_ok=True)

wb = Workbook()

ws = wb.active
ws.title = "Release Summary"

ws["A1"] = "Release Evidence Report"

ws["A3"] = "Repository"
ws["B3"] = os.getenv("GITHUB_REPOSITORY")

ws["A4"] = "Branch"
ws["B4"] = os.getenv("GITHUB_REF_NAME")

ws["A5"] = "Commit SHA"
ws["B5"] = os.getenv("GITHUB_SHA")

ws["A6"] = "Workflow Run ID"
ws["B6"] = os.getenv("GITHUB_RUN_ID")

ws["A7"] = "Triggered By"
ws["B7"] = os.getenv("GITHUB_ACTOR")

ws["A8"] = "Generated On"
ws["B8"] = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")

output = "reports/ReleaseEvidence.xlsx"

wb.save(output)

logger.info("Evidence report generated: %s", output)
